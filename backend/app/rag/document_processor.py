import asyncio
import logging
import os

from langchain_core.documents import Document as LCDocument
from slugify import slugify

from app.core.config import settings

logger = logging.getLogger("campus_ai.rag.document_processor")


class DocumentProcessor:
    def __init__(self):
        self.chunk_size = settings.RAG_CHUNK_SIZE
        self.chunk_overlap = settings.RAG_CHUNK_OVERLAP

    def get_safe_filename(self, original: str) -> str:
        name, ext = os.path.splitext(original)
        return f"{slugify(name) or 'document'}{ext.lower()}"

    def validate_file(self, filename: str, file_size: int) -> tuple[bool, str]:
        allowed_extensions = {".pdf", ".docx", ".txt"}
        _, ext = os.path.splitext(filename.lower())
        if ext not in allowed_extensions:
            return False, "Unsupported file type. Upload a PDF, DOCX, or TXT file."
        if file_size <= 0:
            return False, "The selected file is empty."
        if file_size > settings.MAX_FILE_SIZE_MB * 1024 * 1024:
            return False, f"File size exceeds the {settings.MAX_FILE_SIZE_MB}MB limit."
        return True, ""

    async def load_document(self, file_path: str, file_type: str) -> list[LCDocument]:
        def load_sync() -> list[LCDocument]:
            kind = file_type.lower()
            if kind in {"pdf", "application/pdf"} or file_path.lower().endswith(".pdf"):
                from pypdf import PdfReader
                reader = PdfReader(file_path)
                return [LCDocument(page_content=page.extract_text() or "", metadata={"page": index + 1}) for index, page in enumerate(reader.pages)]
            if kind in {"docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document"} or file_path.lower().endswith(".docx"):
                from docx import Document as DocxDocument
                document = DocxDocument(file_path)
                lines = [paragraph.text for paragraph in document.paragraphs]
                for table in document.tables:
                    lines.extend(" | ".join(cell.text for cell in row.cells) for row in table.rows)
                return [LCDocument(page_content="\n".join(lines), metadata={})]
            if kind in {"txt", "text/plain"} or file_path.lower().endswith(".txt"):
                with open(file_path, "rb") as handle:
                    raw = handle.read()
                try:
                    text = raw.decode("utf-8-sig")
                except UnicodeDecodeError:
                    import chardet
                    encoding = chardet.detect(raw).get("encoding") or "latin-1"
                    text = raw.decode(encoding, errors="replace")
                return [LCDocument(page_content=text, metadata={})]
            raise ValueError("Unsupported file type.")

        try:
            # PyPDF can deadlock in a worker thread on some Windows/Python builds;
            # indexing already runs as a background task, so load synchronously.
            return load_sync()
        except Exception:
            logger.exception("Could not extract document text")
            raise

    def split_documents(self, docs: list[LCDocument]) -> list[LCDocument]:
        chunks: list[LCDocument] = []
        step = max(1, self.chunk_size - self.chunk_overlap)
        for document in docs:
            text = document.page_content.strip()
            for start in range(0, len(text), step):
                content = text[start:start + self.chunk_size].strip()
                if content:
                    chunks.append(LCDocument(page_content=content, metadata=dict(document.metadata)))
                if start + self.chunk_size >= len(text):
                    break
        return chunks

    async def process_document(self, file_path: str, document_id: str, original_filename: str, file_type: str, user_id: str = "") -> list[LCDocument]:
        docs = await self.load_document(file_path, file_type)
        docs = [doc for doc in docs if doc.page_content and doc.page_content.strip()]
        if not docs:
            raise ValueError("No readable text was found in this document.")
        chunks = [chunk for chunk in self.split_documents(docs) if chunk.page_content.strip()]
        for index, chunk in enumerate(chunks):
            chunk.metadata.update({
                "document_id": document_id,
                "document_name": original_filename,
                "file_type": file_type,
                "chunk_index": index,
                "user_id": user_id,
            })
        logger.info("Processed document %s into %s chunks", document_id, len(chunks))
        return chunks
